"""Tests for mkdocs-nav assembly and the template-aware DOCX writer."""

import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

from mdup import convert, convert_mkdocs, mkdocs, render
from mdup.cli import main

MKDOCS_YML = """\
site_name: Demo Report
site_description: A demo subtitle
docs_dir: docs
markdown_extensions:
  - pymdownx.superfences:
      custom_fences:
        - name: mermaid
          format: !!python/name:pymdownx.superfences.fence_code_format
extra:
  mdup:
    status: Draft
nav:
  - Home: index.md
  - Part one:
      - Chapter A: one/a.md
      - Deeper:
          - Chapter B: one/b.md
  - External: https://example.com
  - Annexes:
      - Extra: annex/x.md
"""


@pytest.fixture
def project(tmp_path):
    docs = tmp_path / "docs"
    (docs / "one").mkdir(parents=True)
    (docs / "annex").mkdir()
    (docs / "index.md").write_text("# Welcome\n\nSee [chapter B](one/b.md#b-detail).\n")
    (docs / "one" / "a.md").write_text(
        "# Chapter A\n\nText with a note.[^1] Link to [hidden](../hidden.md).\n\n"
        "[^1]: The footnote.\n"
    )
    (docs / "one" / "b.md").write_text("# Chapter B\n\n## B detail\n\nBody.[^1]\n\n[^1]: Other note.\n")
    (docs / "annex" / "x.md").write_text("No heading here, just text.\n")
    (docs / "hidden.md").write_text("# Hidden\n")
    cfg = tmp_path / "mkdocs.yml"
    cfg.write_text(MKDOCS_YML)
    return cfg


def _assemble(cfg, style="part", tmp=None):
    proj = mkdocs.load(cfg)
    return mkdocs.assemble(
        proj, lambda p: render.parse(p.read_text(), work_dir=tmp)[0], section_style=style
    )


def _make_template(path: Path, body_text="{{TITLE}} / {{STATUS}} / {{MISSING}}") -> Path:
    doc = Document()
    doc.add_paragraph(body_text)
    doc.styles.add_style("Brand Table", WD_STYLE_TYPE.TABLE)
    doc.save(path)
    return path


# ------------------------------------------------------------------ rendering


def test_heading_attr_list_and_mkdocs_slugs(tmp_path):
    body, heads = render.parse(
        "# Intro {.part}\n\n## AMECO2 — the successor\n\n### Custom {#my-id .appendix}\n",
        work_dir=tmp_path,
    )
    assert '<h1 id="intro" class="part">Intro</h1>' in body
    assert 'id="ameco2-the-successor"' in body       # mkdocs-style: no "--"
    assert '<h3 id="my-id" class="appendix">Custom</h3>' in body
    assert "{" not in body


def test_shift_and_namespace_helpers():
    body = '<h1 id="a">A</h1><h6 id="z">Z</h6><a href="#a">x</a>'
    assert "<h3" in render.shift_headings(body, 2)
    assert "<h6" in render.shift_headings(body, 2)      # capped
    ns = render.namespace_ids(body, "p1-")
    assert 'id="p1-a"' in ns and 'href="#p1-a"' in ns


# ------------------------------------------------------------------ mkdocs


def test_mkdocs_load_tolerates_python_tags_and_reads_meta(project):
    proj = mkdocs.load(project)
    kinds = [(e.kind, e.title, e.depth) for e in proj.entries]
    assert kinds == [
        ("page", "Home", 0),
        ("section", "Part one", 0),
        ("page", "Chapter A", 1),
        ("section", "Deeper", 1),
        ("page", "Chapter B", 2),
        ("section", "Annexes", 0),
        ("page", "Extra", 1),
    ]                                                   # external link skipped
    assert proj.meta == {"title": "Demo Report", "subtitle": "A demo subtitle",
                         "status": "Draft"}
    assert proj.entries[-1].appendix


def test_mkdocs_assemble_part_style(project, tmp_path):
    html = _assemble(project, "part", tmp_path)
    assert '<h1 id="s1-part-one" class="part">Part one</h1>' in html
    assert '<h1 id="p2-chapter-a">Chapter A</h1>' in html      # chapter keeps its level
    assert '<h1 id="s3-deeper">Deeper</h1>' in html            # nested section = h1
    assert '<h2 id="p4-chapter-b">Chapter B</h2>' in html      # its page is demoted
    assert '<h1 id="p6-extra" class="appendix">Extra</h1>' in html  # nav title + annex
    assert 'href="#p4-b-detail"' in html                       # cross-page link rewritten
    assert "hidden.md" not in html and "hidden" in html        # out-of-nav link unwrapped
    assert 'id="p2-fn1"' in html and 'id="p4-fn1"' in html     # footnotes kept apart


def test_mkdocs_assemble_heading_style(project, tmp_path):
    html = _assemble(project, "heading", tmp_path)
    assert '<h1 id="s1-part-one">Part one</h1>' in html
    assert '<h2 id="p2-chapter-a">Chapter A</h2>' in html
    assert '<h3 id="p4-chapter-b">Chapter B</h3>' in html
    assert "appendix" not in html                               # only applies at h1


def test_mkdocs_without_nav_walks_tree(project):
    project.write_text("site_name: Tree\n")
    titles = [(e.kind, e.title or e.path.name) for e in mkdocs.load(project).entries]
    assert titles[0] == ("page", "index.md")
    assert ("section", "One") in titles and ("section", "Annex") in titles


def test_mkdocs_bad_config(tmp_path):
    with pytest.raises(FileNotFoundError):
        mkdocs.load(tmp_path / "nope.yml")
    (tmp_path / "docs").mkdir()
    cfg = tmp_path / "mkdocs.yml"
    cfg.write_text("site_name: X\nnav:\n  - Missing: missing.md\n")
    with pytest.raises(ValueError, match="no Markdown pages"):
        mkdocs.load(cfg)


# ------------------------------------------------------------------ template


def test_template_docx_from_mkdocs(project, tmp_path, caplog):
    tpl = _make_template(tmp_path / "brand.docx")
    out = convert_mkdocs(project, formats=["docx"], output=tmp_path / "out",
                         template=tpl, meta={"status": "Final"})
    assert out == [tmp_path / "out" / "Demo_Report.docx"]
    doc = Document(out[0])
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Demo Report / Final / "        # CLI meta beats extra.mdup
    assert "{{MISSING}}" in caplog.text

    styles = {p.text: p.style.name for p in doc.paragraphs}
    assert styles["Part one"] == "Heading 1"            # no "Part Title": falls back
    assert styles["Chapter B"] == "Heading 2"
    assert styles["B detail"] == "Heading 3"
    # Content starts on a new page after the template's own body.
    assert doc.paragraphs[1].paragraph_format.page_break_before

    # Internal link points at a bookmark that exists.
    body = doc.element.body
    anchors = {h.get(qn("w:anchor")) for h in body.iter(qn("w:hyperlink"))}
    marks = {b.get(qn("w:name")) for b in body.iter(qn("w:bookmarkStart"))}
    assert anchors and anchors <= marks

    # Real Word footnotes, one per reference.
    with zipfile.ZipFile(out[0]) as z:
        assert "word/footnotes.xml" in z.namelist()
        fn_xml = z.read("word/footnotes.xml").decode()
    assert "The footnote." in fn_xml and "Other note." in fn_xml
    assert len(list(body.iter(qn("w:footnoteReference")))) == 2


def test_template_styles_lists_code_tables(tmp_path):
    tpl = _make_template(tmp_path / "t.docx", body_text="{{CONTENT}}")
    md = tmp_path / "doc.md"
    md.write_text(
        "---\ntitle: FM Title\n---\n# Doc\n\n1. one\n2. two\n   - nested\n\n"
        "- a\n- b\n\n```python\nx = 1\n```\n\n> [!WARNING]\n> Careful.\n\n"
        "| A | B |\n|---|--:|\n| x | 2 |\n\n![Arch](missing.png)\n"
    )
    out = convert([md], formats=["docx"], output=tmp_path / "doc.docx", template=tpl)[0]
    doc = Document(out)
    paras = doc.paragraphs
    assert paras[0].text == "Doc"                      # {{CONTENT}} anchor replaced
    assert not paras[0].paragraph_format.page_break_before  # empty template: no break
    by_text = {p.text: p for p in paras}
    one, two, nested = by_text["one"], by_text["two"], by_text["nested"]
    assert one.style.name == "List Number"
    num = lambda p: p._p.pPr.numPr
    assert num(one).numId.val == num(two).numId.val
    # Stock Word template: single-level numbering, so depth comes from the style.
    assert nested.style.name == "List Bullet 2"
    # A second ordered list would restart; bullets keep the style's numbering.
    assert by_text["a"].style.name == "List Bullet"
    code = by_text["x = 1"]
    assert code.runs[0].font.name == "Courier New"     # no code style: mono fallback
    assert by_text["Careful."].style.name in ("Intense Quote", "Quote")
    table = doc.tables[0]
    assert table.style.name == "Brand Table"           # template's own table style
    assert table.cell(1, 1).paragraphs[0].alignment == 2  # right-aligned column
    assert any(p.text.startswith("Figure 1. Arch") for p in paras)


def test_template_placeholders_from_front_matter(tmp_path):
    tpl = _make_template(tmp_path / "t.docx", body_text="{{TITLE}}|{{AUTHOR}}")
    md = tmp_path / "doc.md"
    md.write_text("---\ntitle: FM Title\nauthor: Someone\n---\nBody text.\n")
    out = convert([md], formats=["docx"], output=tmp_path / "doc.docx", template=tpl)[0]
    doc = Document(out)
    assert doc.paragraphs[0].text == "FM Title|Someone"
    assert doc.core_properties.title == "FM Title"


def test_template_missing_file(tmp_path):
    md = tmp_path / "doc.md"
    md.write_text("# X\n")
    with pytest.raises(FileNotFoundError):
        convert([md], formats=["docx"], template=tmp_path / "nope.docx")


# ------------------------------------------------------------------ CLI


def test_cli_mkdocs_with_template(project, tmp_path, capsys):
    tpl = _make_template(tmp_path / "brand.docx")
    rc = main(["--mkdocs", str(project), "-f", "docx", "--template", str(tpl),
               "--meta", "status=Approved", "-o", str(tmp_path / "o") + "/"])
    assert rc == 0
    out = Path(capsys.readouterr().out.strip())
    assert out.name == "Demo_Report.docx"
    assert Document(out).paragraphs[0].text.startswith("Demo Report / Approved")


def test_cli_rejects_inputs_with_mkdocs(project, tmp_path):
    md = tmp_path / "a.md"
    md.write_text("# A\n")
    with pytest.raises(SystemExit):
        main([str(md), "--mkdocs", str(project)])
    with pytest.raises(SystemExit):
        main([])
    with pytest.raises(SystemExit):
        main([str(md), "--meta", "novalue"])


BRAND = Path(__file__).parents[1] / "templates" / "word-template" / "MeaningfulData_Template.docx"


@pytest.mark.skipif(not BRAND.is_file(), reason="Meaningful Data template not present")
def test_meaningfuldata_template_roles(tmp_path):
    """The bundle's own example.md maps onto the template's custom styles."""
    md = BRAND.parent / "example.md"
    out = convert([md], formats=["docx"], output=tmp_path / "ex.docx", template=BRAND)[0]
    doc = Document(out)
    styles = {p.style.name for p in doc.paragraphs}
    assert {"Part Title", "Appendix Heading", "Callout", "Callout Note",
            "Callout Warning", "Code Block", "Caption", "List Bullet",
            "List Number"} <= styles
    assert doc.tables[-1].style.name == "Meaningful Data Table"
    assert "{{" not in "".join(p.text for p in doc.paragraphs)
    assert doc.paragraphs[1].text == "AMECO2 Consultancy"   # cover title from front matter
