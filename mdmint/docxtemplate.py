"""HTML fragment → DOCX built on a Word template (``--template``).

The default DOCX path (htmldocx) writes into a blank document and cannot use a
template's styles: it hard-codes fonts, fakes list levels with indents and drops
blockquotes, footnotes and internal links. For branded output we instead walk the
same HTML fragment ourselves (BeautifulSoup) and emit paragraphs that carry the
*template's* named styles, so all look-and-feel comes from the template:

* the template body (cover, contents page…) is kept, ``{{PLACEHOLDER}}`` text in
  body, headers and footers is filled from document metadata, and content is
  appended at the end — or where a ``{{CONTENT}}`` paragraph sits;
* each element maps to a style *role* resolved against the template with
  fallbacks (``Code Block`` → ``HTML Preformatted`` → …), so any template works
  and richer templates simply look better. ``styles`` overrides the mapping;
* lists use real Word numbering (levels, restart per list), footnotes become real
  Word footnotes, headings get bookmarks so in-document links work, figures and
  tables get numbered captions.

Style names that the Meaningful Data template defines (``Part Title``,
``Appendix Heading``, ``Callout Note``…) are the first candidates; the markdown
conventions are those of ``md_to_docx.py`` from that template bundle.
"""

from __future__ import annotations

import base64
import datetime
import io
import logging
import os
import re
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup, Comment, NavigableString, Tag
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.parts.story import StoryPart
from docx.shared import Emu, Pt
from docx.text.paragraph import Paragraph

from . import admonitions

log = logging.getLogger("mdmint")

# role -> candidate style names, first match in the template wins.
_ROLES = {
    "part": ["Part Title"],
    "appendix": ["Appendix Heading"],
    "code": ["Code Block", "Source Code", "HTML Preformatted", "Macro Text"],
    "inline_code": ["Inline Code", "Verbatim Char", "HTML Code"],
    "quote": ["Callout", "Quote", "Block Text"],
    "note": ["Callout Note", "Callout", "Intense Quote", "Quote"],
    "warning": ["Callout Warning", "Callout", "Intense Quote", "Quote"],
    "figure": ["Figure", "Captioned Figure"],
    "caption": ["Caption"],
    "table_heading": ["Table Heading"],
    "table_text": ["Table Text"],
    "list_bullet": ["List Bullet"],
    "list_number": ["List Number"],
    "list_continue": ["List Continue", "List Paragraph"],
    "hyperlink": ["Hyperlink"],
    "toc_heading": ["TOC Heading"],
    "footnote_text": ["Footnote Text", "Reference"],
    "footnote_ref": ["Footnote Reference", "Reference Mark"],
}
_WARNING_TYPES = {"warning", "failure", "danger", "bug"}
_GH_MARKER = re.compile(r"^\s*\[!(\w+)\][ \t]*\n?")
_GH_TYPES = {"NOTE": "note", "TIP": "note", "IMPORTANT": "note",
             "WARNING": "warning", "CAUTION": "warning"}
_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
          "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
_BLOCK_TAGS = {"p", "h1", "h2", "h3", "h4", "h5", "h6", "ul", "ol", "pre",
               "blockquote", "table", "div", "nav", "hr", "section", "figure",
               "details", "summary", "dl", "center", "article", "aside"}
_PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}")
_CONTENT_MARKERS = {"CONTENT", "BODY"}
_CONTENT_ANCHOR = re.compile(r"\s*\{\{\s*(CONTENT|BODY)\s*\}\}\s*", re.IGNORECASE)
_LEN = re.compile(r"([\d.]+)\s*(pt|px|in|cm|mm)?")
_MONO = "Courier New"
_WORD_TABLE_STYLE = re.compile(
    r"^(normal table|table normal|table grid|table \w+ \d|(light|medium|dark|colou?rful)"
    r" |plain table|grid table|list table|table (simple|classic|colorful|columns|contemporary"
    r"|elegant|list|professional|subtle|theme|web|3d))", re.IGNORECASE)
_LIST_INDENT_PT = 18


# ------------------------------------------------------------------ template


def _open_template(path: Path):
    """Open a .docx/.dotx as a python-docx Document (templates need a content-type
    swap, since python-docx only opens the *document* main part type)."""
    if path.suffix.lower() not in (".dotx", ".dotm"):
        return Document(os.fspath(path))
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = re.sub(rb"wordprocessingml\.template\.main\+xml",
                              b"wordprocessingml.document.main+xml", data)
                data = data.replace(b"application/vnd.ms-word.template.macroEnabledTemplate.main+xml",
                                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml")
            zout.writestr(item, data)
    buf.seek(0)
    return Document(buf)


def _story_elements(doc):
    """Root elements of the body and of every header/footer part."""
    yield doc.element.body
    for rel in doc.part.rels.values():
        if rel.reltype in (RT.HEADER, RT.FOOTER) and not rel.is_external:
            element = getattr(rel.target_part, "element", None)
            if element is not None:
                yield element


def fill_placeholders(doc, meta: dict) -> list:
    """Replace ``{{KEY}}`` text with ``meta['key']`` in body, headers and footers.

    Returns the placeholder names found that had no value (they are blanked).
    ``{{CONTENT}}`` / ``{{BODY}}`` are content anchors and left alone.
    """
    values = {str(k).upper(): "" if v is None else str(v) for k, v in meta.items()}
    missing: list[str] = []

    def sub(m: re.Match) -> str:
        key = m.group(1).upper()
        if key in _CONTENT_MARKERS:
            return m.group(0)
        if key not in values and key not in missing:
            missing.append(key)
        return values.get(key, "")

    for root in _story_elements(doc):
        for p in root.iter(qn("w:p")):
            texts = [t for t in p.iter(qn("w:t"))]
            if not texts or "{{" not in "".join(t.text or "" for t in texts):
                continue
            for t in texts:  # common case: each placeholder sits in one run
                if t.text and "{{" in t.text:
                    t.text = _PLACEHOLDER.sub(sub, t.text)
            joined = "".join(t.text or "" for t in texts)
            if _PLACEHOLDER.search(joined):  # split across runs: collapse into the first
                texts[0].text = _PLACEHOLDER.sub(sub, joined)
                for t in texts[1:]:
                    t.text = ""

    cp = doc.core_properties
    if meta.get("title"):
        cp.title = str(meta["title"])
    if meta.get("author"):
        cp.author = str(meta["author"])
    if meta.get("subtitle"):
        cp.subject = str(meta["subtitle"])
    return missing


def _settings_insert(settings, el, before: tuple) -> None:
    """Insert *el* into settings.xml respecting the schema's element order."""
    for name in before:
        found = settings.find(qn(name))
        if found is not None:
            found.addprevious(el)
            return
    settings.append(el)


_AFTER_UPDATE_FIELDS = ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat",
                        "w:docVars", "w:rsids", "w:themeFontLang", "w:clrSchemeMapping",
                        "w:shapeDefaults", "w:decimalSymbol", "w:listSeparator")


def _enable_update_fields(doc) -> None:
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        _settings_insert(settings, el, _AFTER_UPDATE_FIELDS)


# ------------------------------------------------------------------ styles


class _Styles:
    """Resolve style roles against the template, with user overrides."""

    def __init__(self, doc, overrides: dict | None):
        self.doc = doc
        self.by_name = {}
        for s in doc.styles:
            try:
                self.by_name.setdefault(s.name.lower(), s)
            except Exception:  # pragma: no cover - malformed style entries
                continue
        self.overrides = {k: v for k, v in (overrides or {}).items() if v}
        self._cache: dict[str, object] = {}

    def named(self, name: str):
        return self.by_name.get(name.lower())

    def get(self, role: str):
        if role in self._cache:
            return self._cache[role]
        style = None
        if role in self.overrides:
            style = self.named(self.overrides[role])
            if style is None:
                log.warning("template: style %r (for %s) not found; using default",
                            self.overrides[role], role)
        if style is None:
            if role.startswith("heading"):
                level = int(role[len("heading"):])
                for lvl in range(level, 0, -1):
                    style = self.named(f"Heading {lvl}")
                    if style is not None:
                        break
            elif role == "table":
                # The template's own table style: the first one that isn't one of
                # Word's stock styles (which templates often carry unused).
                style = next((s for s in self.doc.styles
                              if s.type == WD_STYLE_TYPE.TABLE
                              and not _WORD_TABLE_STYLE.match(s.name)), None) \
                    or self.named("Table Grid")
            else:
                style = next((self.named(n) for n in _ROLES.get(role, [])
                              if self.named(n) is not None), None)
        self._cache[role] = style
        return style

    def ensure_footnote_styles(self):
        """Footnote styles, created (minimal) if the template has none."""
        styles = self.doc.styles
        if self.get("footnote_text") is None:
            st = styles.add_style("Footnote Text", WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = styles["Normal"]
            st.font.size = Pt(9)
            st.paragraph_format.space_before = Pt(0)
            st.paragraph_format.space_after = Pt(2)
            st.paragraph_format.line_spacing = 1.0
            self._cache["footnote_text"] = st
        if self.get("footnote_ref") is None:
            st = styles.add_style("Footnote Reference", WD_STYLE_TYPE.CHARACTER)
            st.font.superscript = True
            self._cache["footnote_ref"] = st
        return self.get("footnote_text"), self.get("footnote_ref")


# ------------------------------------------------------------------ helpers


def _length_pt(value: str | None) -> float | None:
    if not value:
        return None
    m = _LEN.match(value.strip())
    if not m:
        return None
    num, unit = float(m.group(1)), (m.group(2) or "px")
    return num * {"pt": 1, "px": 0.75, "in": 72, "cm": 72 / 2.54, "mm": 72 / 25.4}[unit]


def _css(tag: Tag) -> dict:
    out = {}
    for decl in (tag.get("style") or "").split(";"):
        if ":" in decl:
            k, v = decl.split(":", 1)
            out[k.strip().lower()] = v.strip()
    return out


def _classes(tag: Tag) -> list:
    cls = tag.get("class") or []
    return cls.split() if isinstance(cls, str) else list(cls)


def _is_blank(node) -> bool:
    return isinstance(node, NavigableString) and not node.strip()


def _trim_el(p_el, left: bool = True) -> None:
    """Strip leading/trailing whitespace from a paragraph's text runs (HTML
    whitespace collapsing leaves stray spaces at the edges)."""
    ts = list(p_el.iter(qn("w:t")))
    for seq, strip in (((ts if left else []), str.lstrip), (reversed(ts), str.rstrip)):
        for t in seq:
            t.text = strip(t.text or "")
            if t.text:
                break


def _text_el(text: str):
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    return t


class _Fmt(dict):
    """Inline formatting state (bold, italic, code…) — copied on nesting."""

    def with_(self, **kw):
        new = _Fmt(self)
        new.update(kw)
        return new


class _PartParent:
    """Minimal parent so a python-docx Paragraph can live in a footnote."""

    def __init__(self, part):
        self.part = part


# ------------------------------------------------------------------ writer


class TemplateWriter:
    def __init__(self, doc, styles: _Styles):
        self.doc = doc
        self.st = styles
        sec = doc.sections[-1]
        self.text_width = int(sec.page_width - sec.left_margin - sec.right_margin)
        self.max_height = int((sec.page_height - sec.top_margin - sec.bottom_margin) * 0.8)
        self.container = doc
        self.reuse_p = None           # empty paragraph to fill first (table cells)
        self.default_style = None     # paragraph style for plain text
        self.quote_style = None
        self.list_stack: list[dict] = []
        self.last_list_p = None
        self.after_table = False
        self.pending_break = False
        self.bookmarks: dict[str, str] = {}
        self._bm_id = 9000
        self.counters: dict[str, int] = {}
        self.footnote_src: dict[str, Tag] = {}
        self._fn_root = None
        self._fn_part = None
        self._fn_next = 1
        self.has_toc = any("TOC" in (t.text or "")
                           for t in doc.element.body.iter(qn("w:instrText")))
        self._bullet_num = self._style_num(self.st.get("list_bullet"))
        self._number_abstract = self._abstract_for(self._style_num(self.st.get("list_number")))

    # ---- document-level --------------------------------------------------

    def write(self, html_body: str) -> None:
        soup = BeautifulSoup(html_body, "html.parser")
        for c in soup.find_all(string=lambda s: isinstance(s, Comment)):
            if "pagebreak" in c.replace("-", "").replace("_", "").lower():
                c.replace_with(soup.new_tag("div", attrs={"class": "mdmint-page-break"}))
            else:
                c.extract()
        self._collect_footnotes(soup)
        for i, h in enumerate(soup.find_all(re.compile(r"^h[1-6]$"))):
            if h.get("id"):
                self.bookmarks.setdefault(h["id"], f"_mdmint{i}")
        self.blocks(soup)

    def _collect_footnotes(self, soup) -> None:
        for sec in soup.find_all("section", class_="footnotes"):
            for li in sec.find_all("li", class_="footnote-item"):
                for back in li.find_all("a", class_="footnote-backref"):
                    back.decompose()
                if li.get("id"):
                    self.footnote_src[li["id"]] = li
            sec.extract()
        for hr in soup.find_all("hr", class_="footnotes-sep"):
            hr.decompose()

    # ---- paragraphs -------------------------------------------------------

    def _para(self, style=None):
        if self.reuse_p is not None:
            p, self.reuse_p = self.reuse_p, None
        else:
            p = self.container.add_paragraph()
        if style is not None:
            p.style = style
        if self.container is self.doc:
            if self.after_table:
                p.paragraph_format.space_before = Pt(12)
                self.after_table = False
            if self.pending_break:
                p.paragraph_format.page_break_before = True
                self.pending_break = False
        return p

    def _text_para(self):
        """A paragraph for running text, styled for the current list/quote context."""
        if self.list_stack:
            ctx, depth = self.list_stack[-1], len(self.list_stack) - 1
            if ctx["first"]:
                ctx["first"] = False
                base = "List Number" if ctx["ordered"] else "List Bullet"
                deeper = self.st.named(f"{base} {depth + 1}") if depth else None
                if ctx["num"] is not None and (depth <= ctx["max_lvl"] or deeper is None):
                    p = self._para(self.st.get("list_number" if ctx["ordered"] else "list_bullet"))
                    self._numpr(p, ctx["num"], min(depth, ctx["max_lvl"]))
                else:
                    # No multi-level numbering: Word's stock templates carry one
                    # single-level style per depth ("List Bullet 2", …) instead.
                    p = self._para(deeper or self.st.named(base))
            else:
                p = self._para(self.st.get("list_continue"))
                p.paragraph_format.left_indent = Pt(_LIST_INDENT_PT * (depth + 1))
            self.last_list_p = p
            return p
        return self._para(self.quote_style or self.default_style)

    @staticmethod
    def _trim(p) -> None:
        _trim_el(p._p)

    # ---- numbering --------------------------------------------------------

    def _style_num(self, style):
        if style is None:
            return None
        ppr = style.element.find(qn("w:pPr"))
        numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
        numid = numpr.find(qn("w:numId")) if numpr is not None else None
        return int(numid.get(qn("w:val"))) if numid is not None else None

    def _numbering(self):
        return self.doc.part.numbering_part.element

    def _abstract_for(self, num_id):
        if num_id is None:
            return None
        for n in self._numbering().findall(qn("w:num")):
            if n.get(qn("w:numId")) == str(num_id):
                a = n.find(qn("w:abstractNumId"))
                return a.get(qn("w:val")) if a is not None else None
        return None

    def _max_level(self, num_id) -> int:
        abstract = self._abstract_for(num_id)
        for a in self._numbering().findall(qn("w:abstractNum")):
            if a.get(qn("w:abstractNumId")) == abstract:
                return max(0, len(a.findall(qn("w:lvl"))) - 1)
        return 0

    def _new_num(self, start: int):
        """A fresh numbering instance so each ordered list restarts at *start*."""
        if self._number_abstract is None:
            return None
        numbering = self._numbering()
        num_id = max((int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))),
                     default=0) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), self._number_abstract)
        num.append(a)
        for lvl in range(9):
            o = OxmlElement("w:lvlOverride")
            o.set(qn("w:ilvl"), str(lvl))
            s = OxmlElement("w:startOverride")
            s.set(qn("w:val"), str(start if lvl == 0 else 1))
            o.append(s)
            num.append(o)
        numbering.append(num)
        return num_id

    @staticmethod
    def _numpr(p, num_id, ilvl) -> None:
        numpr = p._p.get_or_add_pPr().get_or_add_numPr()
        numpr.get_or_add_ilvl().val = ilvl
        numpr.get_or_add_numId().val = num_id

    # ---- blocks -----------------------------------------------------------

    def blocks(self, parent) -> None:
        pending: list = []
        for node in list(parent.children):
            if isinstance(node, Tag) and node.name in _BLOCK_TAGS or \
                    (isinstance(node, Tag) and node.name == "img" and not pending):
                self._flush(pending)
                pending = []
                self.block(node)
            else:
                pending.append(node)
        self._flush(pending)

    def _flush(self, nodes) -> None:
        if all(_is_blank(n) for n in nodes):
            return
        p = self._text_para()
        for n in nodes:
            self.inline(p, n, _Fmt())
        self._trim(p)

    def block(self, node: Tag) -> None:
        name, cls = node.name, _classes(node)
        if re.fullmatch(r"h[1-6]", name):
            self.heading(node, int(name[1]), cls)
        elif name == "p":
            self.paragraph(node)
        elif name in ("ul", "ol"):
            self.list_(node)
        elif name == "pre":
            self.code(node)
        elif name == "blockquote":
            self.blockquote(node)
        elif name == "table":
            if "admonition" in cls:
                self.admonition(node, cls)
            else:
                self.table(node)
        elif name == "hr":
            self.rule()
        elif name == "nav" and "toc" in cls:
            self.toc()
        elif name == "div" and "mdmint-page-break" in cls:
            self.pending_break = True
        elif name == "img":
            self.figure(node)
        else:
            self.blocks(node)

    def heading(self, node, level: int, cls) -> None:
        if "part" in cls:
            style = self.st.get("part") or self.st.get("heading1")
        elif "appendix" in cls:
            style = self.st.get("appendix") or self.st.get("heading1")
        else:
            style = self.st.get(f"heading{level}")
        p = self._para(style)
        if node.get("id") in self.bookmarks:
            self._bookmark(p, self.bookmarks[node["id"]])
        for child in node.children:
            self.inline(p, child, _Fmt())
        self._trim(p)

    def paragraph(self, node: Tag) -> None:
        kids = [c for c in node.children if not _is_blank(c)]
        if kids and all(isinstance(c, Tag) and c.name in ("img", "br") for c in kids):
            for img in (c for c in kids if c.name == "img"):
                self.figure(img, diagram="mermaid-diagram" in _classes(node))
            return
        text = node.get_text()
        nxt = node.find_next_sibling()
        if text.startswith("Table: ") and nxt is not None and nxt.name == "table":
            first = next(node.strings)
            first.replace_with(first.replace("Table: ", "", 1))
            self.caption("Table", node).paragraph_format.keep_with_next = True
            return
        p = self._text_para()
        for child in node.children:
            self.inline(p, child, _Fmt())
        self._trim(p)

    def list_(self, node: Tag) -> None:
        ordered = node.name == "ol"
        if ordered:
            try:
                start = int(node.get("start") or 1)
            except ValueError:
                start = 1
            num = self._new_num(start)
        else:
            num = self._bullet_num
        self.list_stack.append({"ordered": ordered, "num": num, "first": True,
                                "max_lvl": self._max_level(num) if num is not None else 0})
        for li in node.find_all("li", recursive=False):
            self.list_stack[-1]["first"] = True
            self.blocks(li)
        self.list_stack.pop()
        if not self.list_stack and self.last_list_p is not None:
            self.last_list_p.paragraph_format.space_after = Pt(9)
            self.last_list_p = None

    def code(self, node: Tag) -> None:
        style = self.st.get("code")
        p = self._para(style)
        if self.list_stack:
            p.paragraph_format.left_indent = Pt(_LIST_INDENT_PT * len(self.list_stack) + 8)
        lines = node.get_text().rstrip("\n").split("\n")
        for i, line in enumerate(lines):
            run = p.add_run(line)
            if style is None:
                run.font.name = _MONO
            if i < len(lines) - 1:
                run.add_break()

    def _with_quote(self, role: str, content) -> None:
        prev, self.quote_style = self.quote_style, self.st.get(role)
        try:
            self.blocks(content)
        finally:
            self.quote_style = prev

    def blockquote(self, node: Tag) -> None:
        role = "quote"
        first = next((s for s in node.strings if s.strip()), None)
        if first is not None:
            m = _GH_MARKER.match(first)
            if m:
                role = _GH_TYPES.get(m.group(1).upper(), "quote")
                first.replace_with(first[m.end():])
        self._with_quote(role, node)

    def admonition(self, node: Tag, cls) -> None:
        kind = admonitions.resolve(cls[1] if len(cls) > 1 else "note")
        role = "warning" if kind in _WARNING_TYPES else "quote" if kind == "quote" else "note"
        cell = node.find(["td", "th"])
        if cell is not None:
            self._with_quote(role, cell)

    def rule(self) -> None:
        p = self._para()
        bdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" '
                        'w:space="1" w:color="auto"/></w:pBdr>')
        p._p.get_or_add_pPr().insert_element_before(
            bdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind",
            "w:contextualSpacing", "w:jc", "w:outlineLvl", "w:rPr", "w:sectPr", "w:pPrChange")

    def toc(self) -> None:
        if self.has_toc:
            return  # the template already carries a Word TOC field
        self.has_toc = True
        heading = self._para(self.st.get("toc_heading"))
        heading.add_run("Contents").bold = self.st.get("toc_heading") is None
        p = self._para()
        self._field(p, ' TOC \\o "1-3" \\h \\z \\u ', "Update this field (F9) to build the table of contents.")
        _enable_update_fields(self.doc)

    # ---- tables -----------------------------------------------------------

    def table(self, node: Tag) -> None:
        rows = []
        for group in node.find_all(["thead", "tbody", "tfoot", "tr"], recursive=False):
            trs = [group] if group.name == "tr" else group.find_all("tr", recursive=False)
            for tr in trs:
                cells = tr.find_all(["th", "td"], recursive=False)
                header = group.name == "thead" or (cells and all(c.name == "th" for c in cells))
                rows.append((header, cells))
        rows = [r for r in rows if r[1]]
        if not rows:
            return
        ncols = max(len(c) for _, c in rows)
        if self.pending_break and self.container is self.doc:
            self._para()  # carries the page break the table itself can't
        tbl = self.container.add_table(rows=len(rows), cols=ncols)
        style = self.st.get("table")
        if style is not None:
            tbl.style = style
        tblw = tbl._tbl.tblPr.find(qn("w:tblW"))
        if tblw is None:
            tblw = OxmlElement("w:tblW")
            tbl._tbl.tblPr.append(tblw)
        tblw.set(qn("w:type"), "pct")
        tblw.set(qn("w:w"), "5000")
        self._column_widths(tbl, rows, ncols)

        head_style, text_style = self.st.get("table_heading"), self.st.get("table_text")
        saved = (self.container, self.reuse_p, self.default_style, self.quote_style,
                 self.list_stack, self.pending_break)
        try:
            for r_idx, (is_head, cells) in enumerate(rows):
                row = tbl.rows[r_idx]
                if is_head:
                    h = OxmlElement("w:tblHeader")
                    h.set(qn("w:val"), "true")
                    row._tr.get_or_add_trPr().append(h)
                for c_idx in range(ncols):
                    cell = row.cells[c_idx]
                    self.container, self.reuse_p = cell, cell.paragraphs[0]
                    self.default_style = head_style if is_head else text_style
                    self.quote_style, self.list_stack, self.pending_break = None, [], False
                    if c_idx >= len(cells):
                        self.reuse_p.style = self.default_style
                        continue
                    src = cells[c_idx]
                    bold = is_head and head_style is None
                    if any(isinstance(c, Tag) and c.name in _BLOCK_TAGS for c in src.children):
                        self.blocks(src)
                    else:
                        p = self._para(self.default_style)
                        for child in src.children:
                            self.inline(p, child, _Fmt(bold=bold))
                        self._trim(p)
                    align = _css(src).get("text-align")
                    if align in _ALIGN:
                        for p in cell.paragraphs:
                            p.alignment = _ALIGN[align]
        finally:
            (self.container, self.reuse_p, self.default_style, self.quote_style,
             self.list_stack, self.pending_break) = saved
        self.after_table = True

    def _column_widths(self, tbl, rows, ncols) -> None:
        """Share the text width between columns by content.

        Each column first gets enough room for its longest word (so nothing breaks
        mid-word), then the rest is shared in proportion to how much text the
        column holds. If even the minimums don't fit, they are scaled down.
        """
        char = Pt(5.5)          # rough average glyph width at table text size
        padding = Pt(14)        # cell margins
        need = [0] * ncols      # longest word, in characters
        volume = [1.0] * ncols  # typical cell length, capped
        for _, cells in rows:
            for i, c in enumerate(cells):
                text = c.get_text().strip()
                need[i] = max(need[i], max((len(w) for w in text.split()), default=0))
                volume[i] += min(len(text), 80)
        mins = [min(n, 30) * char + padding for n in need]
        spare = self.text_width - sum(mins)
        if spare <= 0:
            widths = [self.text_width * m / sum(mins) for m in mins]
        else:
            widths = [m + spare * v / sum(volume) for m, v in zip(mins, volume)]
        for i, col in enumerate(tbl.columns):
            width = Emu(int(widths[i]))
            col.width = width
            for cell in col.cells:
                cell.width = width

    # ---- figures & captions ----------------------------------------------

    def figure(self, img: Tag, diagram: bool = False) -> None:
        style = self.st.get("figure")
        p = self._para(style)
        if style is None:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._picture(p, img)
        cls = _classes(img)
        caption = img.get("title") or img.get("alt") or ""
        if caption and not diagram and not any(c.startswith("math") for c in cls) \
                and caption != "mermaid diagram":
            self.caption("Figure", text=caption)

    def caption(self, kind: str, node: Tag | None = None, text: str = ""):
        n = self.counters[kind] = self.counters.get(kind, 0) + 1
        p = self._para(self.st.get("caption"))
        self._run(p._p, f"{kind} ", _Fmt(bold=True))
        self._field(p, f" SEQ {kind} \\* ARABIC ", str(n), bold=True)
        self._run(p._p, ". ", _Fmt(bold=True))
        if node is not None:
            for child in node.children:
                self.inline(p, child, _Fmt())
        else:
            self._run(p._p, text, _Fmt())
        self._trim(p)
        return p

    def _picture(self, p, img: Tag) -> None:
        from docx.image.exceptions import UnrecognizedImageError

        src = img.get("src") or ""
        alt = img.get("alt") or os.path.basename(src)
        source = None
        if src.startswith("data:") and ";base64," in src:
            source = io.BytesIO(base64.b64decode(src.split(",", 1)[1]))
        elif src and os.path.isfile(src):
            source = src
        if source is None or src.lower().endswith(".svg"):
            self._run(p._p, f"[image: {alt}]", _Fmt(italic=True))
            return
        try:
            shape = p.add_run().add_picture(source)
        except (UnrecognizedImageError, ValueError, OSError) as exc:
            log.warning("template: cannot embed image %s (%s)", src, exc)
            self._run(p._p, f"[image: {alt}]", _Fmt(italic=True))
            return
        css = _css(img)
        w_pt, h_pt = _length_pt(css.get("width") or img.get("width")), \
            _length_pt(css.get("height") or img.get("height"))
        w, h = shape.width, shape.height
        if w_pt and h_pt:
            w, h = Pt(w_pt), Pt(h_pt)
        elif w_pt:
            w, h = Pt(w_pt), int(h * Pt(w_pt) / w)
        elif h_pt:
            w, h = int(w * Pt(h_pt) / h), Pt(h_pt)
        max_w = self.text_width - (Pt(_LIST_INDENT_PT * len(self.list_stack)) if self.list_stack else 0)
        scale = min(1.0, max_w / w, self.max_height / h) if w and h else 1.0
        shape.width, shape.height = int(w * scale), int(h * scale)

    # ---- inline -----------------------------------------------------------

    def inline(self, p, node, fmt: _Fmt, target=None) -> None:
        target = p._p if target is None else target
        if isinstance(node, Comment):
            return
        if isinstance(node, NavigableString):
            text = re.sub(r"\s+", " ", str(node))
            if text:
                self._run(target, text, fmt)
            return
        name = node.name
        if name in ("strong", "b"):
            fmt = fmt.with_(bold=True)
        elif name in ("em", "i", "cite"):
            fmt = fmt.with_(italic=True)
        elif name in ("s", "del", "strike"):
            fmt = fmt.with_(strike=True)
        elif name in ("u", "ins"):
            fmt = fmt.with_(underline=True)
        elif name in ("code", "kbd", "samp", "tt"):
            fmt = fmt.with_(code=True)
        elif name == "sub":
            fmt = fmt.with_(vert="subscript")
        elif name == "sup":
            if "footnote-ref" in _classes(node):
                self._footnote_ref(p, node, target)
                return
            fmt = fmt.with_(vert="superscript")
        elif name == "br":
            r = OxmlElement("w:r")
            r.append(OxmlElement("w:br"))
            target.append(r)
            return
        elif name == "img":
            self._picture(p, node)
            return
        elif name == "input" and (node.get("type") or "").lower() == "checkbox":
            self._run(target, "☒" if node.has_attr("checked") else "☐", fmt)  # label text brings its own space
            return
        elif name == "a" and node.get("href"):
            link = self._hyperlink(p, target, node["href"])
            if link is not None:
                for child in node.children:
                    self.inline(p, child, fmt.with_(link=True), link)
                return
        for child in node.children:
            self.inline(p, child, fmt, target)

    def _run(self, target, text: str, fmt: _Fmt) -> None:
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rstyle = None
        if fmt.get("code"):
            rstyle = self.st.get("inline_code")
            if rstyle is None:
                fonts = OxmlElement("w:rFonts")
                for a in ("w:ascii", "w:hAnsi", "w:cs"):
                    fonts.set(qn(a), _MONO)
                rpr.append(fonts)
        elif fmt.get("link"):
            rstyle = self.st.get("hyperlink")
        if rstyle is not None:
            rs = OxmlElement("w:rStyle")
            rs.set(qn("w:val"), rstyle.style_id)
            rpr.insert(0, rs)
        for key, tag in (("bold", "w:b"), ("italic", "w:i"), ("strike", "w:strike")):
            if fmt.get(key):
                rpr.append(OxmlElement(tag))
        if fmt.get("link") and rstyle is None:
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "0563C1")
            rpr.append(color)
        if fmt.get("underline") or (fmt.get("link") and rstyle is None):
            u = OxmlElement("w:u")
            u.set(qn("w:val"), "single")
            rpr.append(u)
        if fmt.get("vert"):
            v = OxmlElement("w:vertAlign")
            v.set(qn("w:val"), fmt["vert"])
            rpr.append(v)
        if len(rpr):
            r.append(rpr)
        r.append(_text_el(text))
        target.append(r)

    def _hyperlink(self, p, target, href: str):
        h = OxmlElement("w:hyperlink")
        if href.startswith("#"):
            name = self.bookmarks.get(href[1:])
            if name is None:
                return None  # target isn't a heading we bookmarked: plain text
            h.set(qn("w:anchor"), name)
            h.set(qn("w:history"), "1")
        else:
            h.set(qn("r:id"), p.part.relate_to(href, RT.HYPERLINK, is_external=True))
        target.append(h)
        return h

    def _bookmark(self, p, name: str) -> None:
        self._bm_id += 1
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self._bm_id))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self._bm_id))
        p._p.append(start)
        p._p.append(end)

    def _field(self, p, instr: str, result: str, bold: bool = False) -> None:
        def run(child):
            r = OxmlElement("w:r")
            if bold:
                rpr = OxmlElement("w:rPr")
                rpr.append(OxmlElement("w:b"))
                r.append(rpr)
            r.append(child)
            p._p.append(r)

        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        run(begin)
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        run(it)
        sep = OxmlElement("w:fldChar")
        sep.set(qn("w:fldCharType"), "separate")
        run(sep)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = result
        run(t)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run(end)

    # ---- footnotes --------------------------------------------------------

    def _footnotes(self):
        """The footnotes part's root element, created on first use."""
        if self._fn_root is not None:
            return self._fn_root
        part = None
        for rel in self.doc.part.rels.values():
            if rel.reltype == RT.FOOTNOTES and not rel.is_external:
                part = rel.target_part
                if not hasattr(part, "element"):  # loaded as an opaque blob part
                    part = StoryPart(part.partname, part.content_type,
                                     parse_xml(part.blob), part.package)
                    rel._target = part
                break
        if part is None:
            xml = (f'<w:footnotes {nsdecls("w", "r")}>'
                   '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr>'
                   '<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
                   '</w:pPr><w:r><w:separator/></w:r></w:p></w:footnote>'
                   '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr>'
                   '<w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
                   '</w:pPr><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
                   '</w:footnotes>')
            part = StoryPart(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES,
                             parse_xml(xml), self.doc.part.package)
            self.doc.part.relate_to(part, RT.FOOTNOTES)
            settings = self.doc.settings.element
            if settings.find(qn("w:footnotePr")) is None:
                pr = parse_xml(f'<w:footnotePr {nsdecls("w")}><w:footnote w:id="-1"/>'
                               '<w:footnote w:id="0"/></w:footnotePr>')
                _settings_insert(settings, pr, _AFTER_UPDATE_FIELDS[2:])
        self._fn_part, self._fn_root = part, part.element
        ids = [int(f.get(qn("w:id"))) for f in self._fn_root.findall(qn("w:footnote"))]
        self._fn_next = max(ids + [0]) + 1
        return self._fn_root

    def _footnote_ref(self, p, sup: Tag, target) -> None:
        a = sup.find("a")
        src = self.footnote_src.get((a.get("href") or "")[1:]) if a is not None else None
        if src is None:
            self._run(target, sup.get_text(), _Fmt(vert="superscript"))
            return
        root = self._footnotes()
        text_style, ref_style = self.st.ensure_footnote_styles()
        fid = self._fn_next
        self._fn_next += 1

        def ref_run(tag):
            r = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            rs = OxmlElement("w:rStyle")
            rs.set(qn("w:val"), ref_style.style_id)
            rpr.append(rs)
            r.append(rpr)
            el = OxmlElement(tag)
            if tag == "w:footnoteReference":
                el.set(qn("w:id"), str(fid))
            r.append(el)
            return r

        fn = OxmlElement("w:footnote")
        fn.set(qn("w:id"), str(fid))
        root.append(fn)
        parent = _PartParent(self._fn_part)
        tab = bool(len(text_style.paragraph_format.tab_stops))  # style lays out "1<tab>text"
        blocks = [c for c in src.children if isinstance(c, Tag)] or [src]
        for i, block in enumerate(blocks):
            p_el = OxmlElement("w:p")
            ppr = OxmlElement("w:pPr")
            ps = OxmlElement("w:pStyle")
            ps.set(qn("w:val"), text_style.style_id)
            ppr.append(ps)
            p_el.append(ppr)
            fn.append(p_el)
            fp = Paragraph(p_el, parent)
            if i == 0:
                p_el.append(ref_run("w:footnoteRef"))
                r = OxmlElement("w:r")
                r.append(OxmlElement("w:tab") if tab else _text_el(" "))
                p_el.append(r)
            children = block.children if block.name == "p" else [block]
            for child in children:
                self.inline(fp, child, _Fmt())
            _trim_el(p_el, left=False)
        target.append(ref_run("w:footnoteReference"))


# ------------------------------------------------------------------ entry


def _body_has_content(body, upto=None) -> bool:
    for el in body:
        if el is upto:
            break
        if el.tag == qn("w:tbl") or el.find(".//" + qn("w:drawing")) is not None:
            return True
        if el.tag == qn("w:p") and "".join(t.text or "" for t in el.iter(qn("w:t"))).strip():
            return True
    return False


def to_docx(body: str, out_path, template, meta: dict | None = None,
            styles: dict | None = None) -> None:
    """Write *body* (an HTML fragment) into a copy of *template* at *out_path*."""
    template = Path(template)
    if not template.is_file():
        raise FileNotFoundError(f"DOCX template not found: {template}")
    doc = _open_template(template)
    meta = dict(meta or {})
    meta.setdefault("date", datetime.date.today().strftime("%B %Y"))
    missing = fill_placeholders(doc, meta)
    if missing:
        log.warning("template: no value for %s (left blank; set with --meta KEY=VALUE)",
                    ", ".join("{{%s}}" % m for m in missing))

    body_el = doc.element.body
    marker = next((p for p in body_el.iterchildren(qn("w:p"))
                   if _CONTENT_ANCHOR.fullmatch("".join(t.text or "" for t in p.iter(qn("w:t"))))),
                  None)
    before = set(body_el)  # holding the proxies keeps their identity stable

    writer = TemplateWriter(doc, _Styles(doc, styles))
    writer.pending_break = _body_has_content(body_el, marker)
    writer.write(body)

    if marker is not None:  # move the new content to the {{CONTENT}} anchor
        new = [el for el in body_el if el not in before and el.tag != qn("w:sectPr")]
        for el in new:
            marker.addprevious(el)
        if marker.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None:
            for r in marker.findall(qn("w:r")):
                marker.remove(r)
        else:
            body_el.remove(marker)
    if writer.has_toc:
        _enable_update_fields(doc)
    doc.save(os.fspath(out_path))
